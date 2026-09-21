"""Add a translation after every paragraph of an existing library document
(docx / txt) -- the bilingual layout Rey's curriculum is sent in: an English
paragraph, immediately followed by its Spanish. Added 20Sep2026 for Reference
Hub's "Translate to Spanish" button.

The output is a *copy of the original file* (written to a separate output
folder, never beside the source) with Spanish paragraphs inserted in place,
not a rebuilt document -- so headings, tables, fonts, spacing and
page setup carry over, and each Spanish paragraph copies its English
neighbor's paragraph style (a Spanish heading is a heading). Trade-offs:
mixed formatting inside one paragraph (a bold word mid-sentence) collapses to
the first run's style on the Spanish side, since a translated sentence
doesn't map word-for-word onto the original's runs; and a numbered-list
paragraph's Spanish line is indented but unnumbered, so it doesn't consume a
list number.
"""
import copy
import os
from pathlib import Path
from typing import Callable, Dict, List

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

TRANSLATABLE_SUFFIXES = {".docx", ".txt"}


def _open_permissions(path: Path, mode: int) -> None:
    """The backend container runs as root, so anything it creates under the
    bind-mounted ./data is root-owned on the host -- Rey couldn't delete or
    move his own output from WSL without sudo. Loosen it on what we create.
    Best-effort: a failure here must never fail the translation itself."""
    try:
        os.chmod(path, mode)
    except OSError:
        pass


def _output_path(src: Path, label: str, out_dir: Path) -> Path:
    """{stem}_{label}{ext} in `out_dir`; `_bilingual_2`, `_3`... rather than
    silently overwriting an earlier (possibly hand-corrected) copy -- which
    also covers two same-named sources from different subfolders landing in
    the same flat output folder. Underscore between words, per the project's
    filename convention."""
    out_dir.mkdir(parents=True, exist_ok=True)
    _open_permissions(out_dir, 0o777)
    candidate = out_dir / f"{src.stem}_{label}{src.suffix}"
    n = 2
    while candidate.exists():
        candidate = out_dir / f"{src.stem}_{label}_{n}{src.suffix}"
        n += 1
    return candidate


def _iter_paragraphs(doc):
    """Body paragraphs plus every table cell (recursively -- nested tables)."""
    def from_container(container):
        for p in container.paragraphs:
            yield p
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from from_container(cell)
    yield from from_container(doc)


def _insert_translation_after(p, text: str):
    """Clone `p`'s paragraph properties (style, alignment, spacing), drop its
    content, and put `text` in one run that borrows the first run's font."""
    src_el = p._p
    new_el = copy.deepcopy(src_el)
    for child in list(new_el):
        if child.tag != qn("w:pPr"):
            new_el.remove(child)
    src_el.addnext(new_el)
    new_p = Paragraph(new_el, p._parent)

    pPr = new_el.find(qn("w:pPr"))
    was_list_item = pPr is not None and pPr.find(qn("w:numPr")) is not None
    if was_list_item:
        pPr.remove(pPr.find(qn("w:numPr")))
        new_p.paragraph_format.left_indent = Inches(0.25)

    run = new_p.add_run(text)
    if p.runs:
        rPr = p.runs[0]._r.find(qn("w:rPr"))
        if rPr is not None:
            run._r.insert(0, copy.deepcopy(rPr))
    return new_p


def translate_document(
    src: Path,
    label: str,
    translate_fn: Callable[[List[str]], List[str]],
    out_dir: Path,
) -> Dict:
    suffix = src.suffix.lower()
    if suffix not in TRANSLATABLE_SUFFIXES:
        raise ValueError(
            f"Can't translate {suffix or 'this file type'} yet -- only .docx and .txt. "
            "(PDFs are layout-only; export the source as .docx first.)"
        )
    dest = _output_path(src, label, out_dir)

    if suffix == ".txt":
        lines = src.read_text(encoding="utf-8", errors="replace").splitlines()
        idx = [i for i, l in enumerate(lines) if l.strip()]
        translated = dict(zip(idx, translate_fn([lines[i] for i in idx])))
        out = []
        for i, l in enumerate(lines):
            out.append(l)
            if i in translated:
                out.append(translated[i])
        dest.write_text("\n".join(out) + "\n", encoding="utf-8")
        _open_permissions(dest, 0o666)
        return {"path": dest, "translated": len(idx), "skipped": 0}

    doc = Document(str(src))
    paras = [p for p in _iter_paragraphs(doc) if p.text.strip()]
    # Identical strings (repeated headings, "Yes"/"No" cells) translated once.
    unique = list(dict.fromkeys(p.text for p in paras))
    lookup = dict(zip(unique, translate_fn(unique)))

    skipped = 0
    for p in paras:
        new = lookup.get(p.text)
        if not new:
            skipped += 1
            continue
        _insert_translation_after(p, new)
    doc.save(str(dest))
    _open_permissions(dest, 0o666)
    return {"path": dest, "translated": len(paras) - skipped, "skipped": skipped}
